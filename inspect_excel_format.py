import pathlib
import openpyxl
import itertools

p = pathlib.Path('lib') / 'TY A TT.xlsx'
print('path:', p)
print('exists:', p.exists())
if not p.exists():
    raise SystemExit(1)
wb = openpyxl.load_workbook(p, data_only=True)
print('sheets:', wb.sheetnames)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print('sheet:', sheet_name)
    for i, row in enumerate(itertools.islice(ws.iter_rows(values_only=True), 20), 1):
        print(i, row)
    break
